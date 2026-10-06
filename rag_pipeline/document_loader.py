# document_loader.py

import os
from pathlib import Path
from typing import List, Dict, Optional
from dataclasses import dataclass
import logging

import PyPDF2
import pdfplumber
from docx import Document as DocxDocument

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


@dataclass
class LoadedDocument:
    """Represents a loaded document with its content and metadata."""
    content: str
    metadata: Dict
    source: str
    file_type: str
    page_count: int = 0


class DocumentLoader:
    """Loads documents from PDF, DOCX, and TXT files."""

    def __init__(self):
        self.loaders = {
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
            ".doc": self._load_docx,
            ".txt": self._load_txt,
        }

    def load_file(self, file_path: str) -> Optional[LoadedDocument]:
        """Load a single file and return a LoadedDocument."""
        path = Path(file_path)

        if not path.exists():
            logger.error(f"File not found: {file_path}")
            return None

        ext = path.suffix.lower()
        if ext not in self.loaders:
            logger.error(f"Unsupported file type: {ext}")
            return None

        logger.info(f"Loading: {path.name}")
        try:
            return self.loaders[ext](path)
        except Exception as e:
            logger.error(f"Error loading {path.name}: {e}")
            return None

    def load_directory(self, dir_path: str, recursive: bool = True) -> List[LoadedDocument]:
        """Load all supported documents from a directory."""
        documents = []
        path = Path(dir_path)

        if not path.exists():
            logger.error(f"Directory not found: {dir_path}")
            return documents

        # Gather all files
        pattern = "**/*" if recursive else "*"
        files = [
            f for f in path.glob(pattern)
            if f.is_file() and f.suffix.lower() in self.loaders
        ]

        logger.info(f"Found {len(files)} supported files in {dir_path}")

        for file in files:
            doc = self.load_file(str(file))
            if doc:
                documents.append(doc)

        logger.info(f"Successfully loaded {len(documents)} documents")
        return documents

    # ── Private Loaders ──

    def _load_pdf(self, path: Path) -> LoadedDocument:
        """Load a PDF file using pdfplumber (better text extraction)."""
        pages_text = []
        page_count = 0

        try:
            # Try pdfplumber first (better formatting)
            with pdfplumber.open(path) as pdf:
                page_count = len(pdf.pages)
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)
        except Exception:
            # Fallback to PyPDF2
            logger.warning(f"pdfplumber failed for {path.name}, trying PyPDF2")
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                page_count = len(reader.pages)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)

        content = "\n\n".join(pages_text)

        return LoadedDocument(
            content=content,
            metadata={
                "source": str(path),
                "filename": path.name,
                "file_type": "pdf",
                "page_count": page_count,
                "file_size_kb": round(path.stat().st_size / 1024, 2),
            },
            source=str(path),
            file_type="pdf",
            page_count=page_count,
        )

    def _load_docx(self, path: Path) -> LoadedDocument:
        """Load a Word document."""
        doc = DocxDocument(path)

        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        content = "\n\n".join(paragraphs)

        return LoadedDocument(
            content=content,
            metadata={
                "source": str(path),
                "filename": path.name,
                "file_type": "docx",
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
                "file_size_kb": round(path.stat().st_size / 1024, 2),
            },
            source=str(path),
            file_type="docx",
        )

    def _load_txt(self, path: Path) -> LoadedDocument:
        """Load a plain text file."""
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()

        return LoadedDocument(
            content=content,
            metadata={
                "source": str(path),
                "filename": path.name,
                "file_type": "txt",
                "file_size_kb": round(path.stat().st_size / 1024, 2),
            },
            source=str(path),
            file_type="txt",
        )